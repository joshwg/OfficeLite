#!/usr/bin/env python3
"""Regenerates testdata/pydocx.docx with python-docx (pip install python-docx).

Run from this directory: python3 gen_pydocx.py
"""
import io, struct, zlib
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor, Inches


def tiny_png():
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * 4 for _ in range(4))
    def chunk(t, d):
        c = struct.pack(">I", len(d)) + t + d
        return c + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 4, 4, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


d = Document()
d.core_properties.author = "test"
d.add_heading("Fixture Title", 0)
d.add_heading("Heading One", 1)
p = d.add_paragraph("Plain ")
r = p.add_run("bold"); r.bold = True
p.add_run(" and ")
r = p.add_run("italic red"); r.italic = True; r.font.color.rgb = RGBColor(0xFF, 0, 0)
p.add_run(" ")
r = p.add_run("big Arial"); r.font.size = Pt(16); r.font.name = "Arial"
p.add_run("\ttabbed  two spaces")
r = p.add_run("under"); r.underline = True
r = p.add_run("not bold"); r.bold = False
p = d.add_paragraph("Centered", style="Quote"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6); p.paragraph_format.left_indent = Inches(0.5)
p.paragraph_format.first_line_indent = Inches(-0.25)
d.add_paragraph("bullet one", style="List Bullet")
d.add_paragraph("bullet two", style="List Bullet")
d.add_paragraph("number one", style="List Number")
d.add_paragraph("number two", style="List Number")
t = d.add_table(rows=3, cols=3); t.style = "Table Grid"
for i, row in enumerate(t.rows):
    for j, c in enumerate(row.cells):
        c.text = f"r{i}c{j}"
t.cell(0, 0).merge(t.cell(0, 1))
t.cell(1, 2).merge(t.cell(2, 2))
p = d.add_paragraph("before break")
p.add_run().add_break(WD_BREAK.PAGE)
p.add_run("after break")
d.add_picture(io.BytesIO(tiny_png()), width=Inches(1))
d.add_heading("Heading Two", 2)
d.add_paragraph("Last paragraph.")
d.save("pydocx.docx")
